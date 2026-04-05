"""Compare two generated reconciliation workbooks for parity diagnostics."""

import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    """Parse command-line arguments for workbook comparison."""
    parser = argparse.ArgumentParser(
        description=(
            'Compare two attached reconciliation workbooks with '
            'summary and row-level diagnostics.'
        )
    )
    parser.add_argument(
        '--baseline', required=True, help='Path to the baseline workbook.'
    )
    parser.add_argument(
        '--candidate', required=True, help='Path to the candidate workbook.'
    )
    return parser.parse_args()


ERROR_REPORT_SHEET = 'Error Report'


def s(v):
    """Return a stripped string representation for a cell value."""
    return '' if v is None else str(v).strip()


def num(v):
    """Parse a workbook currency value into a float when possible."""
    t = (
        s(v)
        .replace(',', '')
        .replace('$', '')
        .replace('(', '-')
        .replace(')', '')
    )
    if not t:
        return None
    try:
        return float(Decimal(t))
    except (InvalidOperation, ValueError):
        return None


def read_summary(ws):
    """Read the top summary key-value pairs from a worksheet."""
    d = {}
    for r in range(1, min(ws.max_row, 300) + 1):
        k = s(ws.cell(r, 1).value)
        v = s(ws.cell(r, 2).value)
        if k:
            d[k] = v
    return d


def detect_header_row(ws, max_scan=12):
    """Find the most likely header row by counting non-empty cells."""
    best = (-1, 1)
    for r in range(1, min(ws.max_row, max_scan) + 1):
        c = 0
        for col in range(1, min(ws.max_column, 100) + 1):
            if s(ws.cell(r, col).value):
                c += 1
        if c > best[0]:
            best = (c, r)
    return best[1]


def read_table(ws):
    """Read a worksheet into header metadata and row dictionaries."""
    hr = detect_header_row(ws)
    cols = []
    headers = []
    for c in range(1, min(ws.max_column, 100) + 1):
        h = s(ws.cell(hr, c).value)
        if h:
            cols.append(c)
            headers.append(h)

    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        vals = [s(ws.cell(r, c).value) for c in cols]
        if any(vals):
            rows.append(dict(zip(headers, vals)))
    return hr, headers, rows


def get_cat(row):
    """Extract the category-like field from a workbook row."""
    for k in ['Error Category', 'Category', 'Exception Category', 'Status']:
        if row.get(k):
            return row[k]
    return ''


def get_amt(row):
    """Extract the monetary exposure field from a workbook row."""
    for k in ['Amount', 'Dollar Exposure', 'Exposure']:
        if row.get(k):
            return num(row[k]) or 0.0
    return 0.0


def main():
    """Compare workbook summaries, error categories, and row-level deltas."""
    args = parse_args()
    a_path = Path(args.baseline).expanduser().resolve()
    b_path = Path(args.candidate).expanduser().resolve()

    if not a_path.exists() or not b_path.exists():
        print('ERROR: workbook path missing')
        print(f'  A exists: {a_path.exists()} => {a_path}')
        print(f'  B exists: {b_path.exists()} => {b_path}')
        raise SystemExit(1)

    wa = load_workbook(a_path, data_only=True)
    wb = load_workbook(b_path, data_only=True)

    print('A', a_path)
    print('B', b_path)
    print('SHEETS_A', wa.sheetnames)
    print('SHEETS_B', wb.sheetnames)

    if 'Summary' in wa.sheetnames and 'Summary' in wb.sheetnames:
        sa = read_summary(wa['Summary'])
        sb = read_summary(wb['Summary'])
        keys = sorted(set(sa) | set(sb))
        diffs = [
            (k, sa.get(k, ''), sb.get(k, ''))
            for k in keys
            if sa.get(k, '') != sb.get(k, '')
        ]
        print('SUMMARY_DIFF_COUNT', len(diffs))
        for k, av, bv in diffs[:60]:
            print('SUMMARY_DIFF', k, '|A=', av, '|B=', bv)

    if (
        ERROR_REPORT_SHEET in wa.sheetnames
        and ERROR_REPORT_SHEET in wb.sheetnames
    ):
        hra, hda, ra = read_table(wa[ERROR_REPORT_SHEET])
        hrb, hdb, rb = read_table(wb[ERROR_REPORT_SHEET])
        print(
            'ERROR_REPORT A_header', hra, 'A_rows', len(ra), 'A_cols', len(hda)
        )
        print(
            'ERROR_REPORT B_header', hrb, 'B_rows', len(rb), 'B_cols', len(hdb)
        )

        ca = Counter(get_cat(r) for r in ra if get_cat(r))
        cb = Counter(get_cat(r) for r in rb if get_cat(r))
        cats = sorted(set(ca) | set(cb))
        print('CATEGORY_COUNT_DELTAS')
        for c in cats:
            if ca.get(c, 0) != cb.get(c, 0):
                print(
                    'CAT',
                    c,
                    'A',
                    ca.get(c, 0),
                    'B',
                    cb.get(c, 0),
                    'DELTA',
                    cb.get(c, 0) - ca.get(c, 0),
                )

        ea = Counter()
        eb = Counter()
        for r in ra:
            c = get_cat(r)
            if c:
                ea[c] += get_amt(r)
        for r in rb:
            c = get_cat(r)
            if c:
                eb[c] += get_amt(r)
        print('CATEGORY_EXPOSURE_DELTAS')
        for c in cats:
            av = round(ea.get(c, 0.0), 2)
            bv = round(eb.get(c, 0.0), 2)
            if av != bv:
                print('EXP', c, 'A', av, 'B', bv, 'DELTA', round(bv - av, 2))

    # Row-level mapping (known column layout per workbook flavor)
    wsa = wa[ERROR_REPORT_SHEET]
    wsb = wb[ERROR_REPORT_SHEET]
    rows_a = []
    for r in range(2, wsa.max_row + 1):
        source = s(wsa.cell(r, 1).value)
        line = s(wsa.cell(r, 3).value)
        cat = s(wsa.cell(r, 10).value)
        amt = num(wsa.cell(r, 9).value)
        msg = s(wsa.cell(r, 11).value)
        if source or line or cat:
            rows_a.append(
                {
                    'key': (source.lower(), line),
                    'source': source,
                    'line': line,
                    'cat': cat,
                    'amount': amt,
                    'msg': msg,
                }
            )

    rows_b = []
    for r in range(2, wsb.max_row + 1):
        source = s(wsb.cell(r, 3).value)
        line = s(wsb.cell(r, 5).value)
        cat = s(wsb.cell(r, 2).value)
        amt = num(wsb.cell(r, 7).value)
        msg = s(wsb.cell(r, 8).value)
        if source or line or cat:
            rows_b.append(
                {
                    'key': (source.lower(), line),
                    'source': source,
                    'line': line,
                    'cat': cat,
                    'amount': amt,
                    'msg': msg,
                }
            )

    idxa = defaultdict(list)
    idxb = defaultdict(list)
    for x in rows_a:
        idxa[x['key']].append(x)
    for x in rows_b:
        idxb[x['key']].append(x)

    common = set(idxa) & set(idxb)
    only_a = set(idxa) - set(idxb)
    only_b = set(idxb) - set(idxa)

    print(
        'KEY_COVERAGE A',
        len(idxa),
        'B',
        len(idxb),
        'COMMON',
        len(common),
        'ONLY_A',
        len(only_a),
        'ONLY_B',
        len(only_b),
    )

    mismatch = Counter()
    for k in common:
        left = idxa[k][0]['cat']
        right = idxb[k][0]['cat']
        if left != right:
            mismatch[(left, right)] += 1

    print('TOP_MISMATCH_PAIRS')
    for (left, right), count in mismatch.most_common(20):
        print('PAIR', left, '->', right, 'COUNT', count)

    print('SAMPLE_ONLY_B')
    for k in sorted(only_b)[:30]:
        r = idxb[k][0]
        print(
            'ONLY_B',
            r['source'],
            'line',
            r['line'],
            'cat',
            r['cat'],
            'amt',
            r['amount'],
            'msg',
            r['msg'][:110],
        )


if __name__ == '__main__':
    main()
