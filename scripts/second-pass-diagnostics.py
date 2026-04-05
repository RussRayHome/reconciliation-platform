"""Run second-pass diagnostics against two reconciliation workbooks."""

import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    """Parse command-line arguments for the second-pass comparison."""
    parser = argparse.ArgumentParser(
        description=(
            'Run second-pass diagnostics against two workbooks '
            'using the expected Error Report layouts.'
        )
    )
    parser.add_argument(
        '--baseline', required=True, help='Path to the baseline workbook.'
    )
    parser.add_argument(
        '--candidate', required=True, help='Path to the candidate workbook.'
    )
    return parser.parse_args()


def s(v):
    """Return a stripped string representation for a worksheet value."""
    return '' if v is None else str(v).strip()


def num(v):
    """Parse a worksheet amount cell into a float when possible."""
    t = (
        s(v)
        .replace(',', '')
        .replace('$', '')
        .replace('(', '-')
        .replace(')', '')
    )
    if not t:
        return None
    try:
        return float(Decimal(t))
    except (InvalidOperation, ValueError):
        return None


def read_a(path: Path):
    """Read rows from the baseline workbook error report layout."""
    wb = load_workbook(path, data_only=True)
    ws = wb['Error Report']
    rows = []
    for r in range(2, ws.max_row + 1):
        source = s(ws.cell(r, 1).value)
        line = s(ws.cell(r, 3).value)
        date = s(ws.cell(r, 4).value)
        amt = num(ws.cell(r, 9).value)
        cat = s(ws.cell(r, 10).value)
        msg = s(ws.cell(r, 11).value)
        if not source and not line and not cat:
            continue
        key = (source.lower(), line)
        rows.append({
            'key': key,
            'source': source,
            'line': line,
            'date': date,
            'amount': amt,
            'cat': cat,
            'msg': msg,
            'row': r,
        })
    return rows


def read_b(path: Path):
    """Read rows from the candidate workbook error report layout."""
    wb = load_workbook(path, data_only=True)
    ws = wb['Error Report']
    rows = []
    for r in range(2, ws.max_row + 1):
        source = s(ws.cell(r, 3).value)
        line = s(ws.cell(r, 5).value)
        date = s(ws.cell(r, 6).value)
        amt = num(ws.cell(r, 7).value)
        cat = s(ws.cell(r, 2).value)
        msg = s(ws.cell(r, 8).value)
        if not source and not line and not cat:
            continue
        key = (source.lower(), line)
        rows.append({
            'key': key,
            'source': source,
            'line': line,
            'date': date,
            'amount': amt,
            'cat': cat,
            'msg': msg,
            'row': r,
        })
    return rows


def main():
    """Print category, coverage, and exposure diagnostics.

    The report compares both workbook layouts and their category totals.
    """
    args = parse_args()
    a_path = Path(args.baseline).expanduser().resolve()
    b_path = Path(args.candidate).expanduser().resolve()

    if not a_path.exists() or not b_path.exists():
        print('ERROR: workbook path missing')
        print(f'  A exists: {a_path.exists()} => {a_path}')
        print(f'  B exists: {b_path.exists()} => {b_path}')
        raise SystemExit(1)

    a = read_a(a_path)
    b = read_b(b_path)

    idx_a = defaultdict(list)
    idx_b = defaultdict(list)
    for x in a:
        idx_a[x['key']].append(x)
    for x in b:
        idx_b[x['key']].append(x)

    common = set(idx_a).intersection(idx_b)
    only_a = set(idx_a) - set(idx_b)
    only_b = set(idx_b) - set(idx_a)

    print('KEY COVERAGE')
    print(
        f'A keys: {len(idx_a)} B keys: {len(idx_b)} Common: {len(common)} '
        f'OnlyA: {len(only_a)} OnlyB: {len(only_b)}'
    )

    cat_mismatch = Counter()
    cat_match = Counter()
    examples = []

    for k in sorted(common):
        ra = idx_a[k][0]
        rb = idx_b[k][0]
        ca, cb = ra['cat'], rb['cat']
        if ca == cb:
            cat_match[ca] += 1
        else:
            cat_mismatch[(ca, cb)] += 1
            da = None
            if ra['amount'] is not None and rb['amount'] is not None:
                da = round(rb['amount'] - ra['amount'], 2)
            examples.append((abs(ra['amount'] or 0), ra, rb, da))

    print('\nTOP CATEGORY MISMATCH PAIRS (A -> B)')
    for (ca, cb), n in cat_mismatch.most_common(15):
        print(f'{n:4d} | {ca} -> {cb}')

    print('\nTOP CATEGORY MATCHES')
    for c, n in cat_match.most_common(10):
        print(f'{n:4d} | {c}')

    print(f'\nONLY IN A (baseline-only keys): {len(only_a)}')
    print(f'ONLY IN B (app-only keys): {len(only_b)}')

    print('\nSAMPLE ONLY-IN-B KEYS (first 30)')
    for k in sorted(only_b)[:30]:
        r = idx_b[k][0]
        print(
            f"{r['source']} line {r['line']} | cat={r['cat']} | "
            f"amount={r['amount']} | "
            f"msg={r['msg'][:90]}"
        )

    print('\nTOP 50 HIGH-IMPACT MISMATCH EXAMPLES (by abs baseline amount)')
    examples.sort(reverse=True, key=lambda t: t[0])
    for _, ra, rb, da in examples[:50]:
        print(
            f"{ra['source']} line {ra['line']} | A:{ra['cat']} "
            f"amt={ra['amount']} -> "
            f"B:{rb['cat']} amt={rb['amount']} | dAmt={da} | "
            f"Amsg={ra['msg'][:70]} | Bmsg={rb['msg'][:70]}"
        )

    sum_a = Counter()
    sum_b = Counter()
    for r in a:
        if r['amount'] is not None:
            sum_a[r['cat']] += r['amount']
    for r in b:
        if r['amount'] is not None:
            sum_b[r['cat']] += r['amount']

    all_cats = sorted(set(sum_a) | set(sum_b))
    print('\nCATEGORY EXPOSURE DELTA (B - A)')
    for c in all_cats:
        da = round(sum_b[c] - sum_a[c], 2)
        if abs(da) > 0.01:
            print(f"{c}: {da:+,.2f} (A={sum_a[c]:,.2f} B={sum_b[c]:,.2f})")


if __name__ == '__main__':
    main()
