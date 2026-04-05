"""Compare reconciliation workbooks at the summary and error-report levels."""

import argparse
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def parse_args():
    """Parse command-line arguments for workbook parity checks."""
    parser = argparse.ArgumentParser(
        description=(
            'Compare two reconciliation workbooks at the summary '
            'and error-report level.'
        )
    )
    parser.add_argument(
        '--baseline', required=True, help='Path to the baseline workbook.'
    )
    parser.add_argument(
        '--candidate', required=True, help='Path to the candidate workbook.'
    )
    return parser.parse_args()


ERROR_REPORT_SHEET = 'Error Report'
REVIEW_SHEETS = [
    'Mapping and Assumptions',
    'Probable Matches Review',
    'Reviewer Observations',
]


def sval(v):
    """Normalize worksheet values into stable string representations."""
    if v is None:
        return ''
    if isinstance(v, float):
        return ('%.10f' % v).rstrip('0').rstrip('.')
    return str(v).strip()


def read_summary(ws):
    """Read summary rows from a workbook summary worksheet."""
    d = {}
    max_r = min(ws.max_row, 250)
    for r in range(1, max_r + 1):
        k = sval(ws.cell(r, 1).value)
        v = sval(ws.cell(r, 2).value)
        if k:
            d[k] = v
    return d


def detect_header_row(ws, max_scan=12):
    """Detect the best table header row within the scan range."""
    best_count = -1
    best_row = 1
    for r in range(1, min(ws.max_row, max_scan) + 1):
        non_empty = 0
        for c in range(1, min(ws.max_column, 80) + 1):
            if sval(ws.cell(r, c).value):
                non_empty += 1
        if non_empty > best_count:
            best_count = non_empty
            best_row = r
    return best_row


def read_table(ws):
    """Read a table worksheet into header metadata and row dictionaries."""
    hr = detect_header_row(ws)
    idx = []
    headers = []
    for c in range(1, min(ws.max_column, 80) + 1):
        h = sval(ws.cell(hr, c).value)
        if h:
            idx.append(c)
            headers.append(h)

    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        vals = [sval(ws.cell(r, c).value) for c in idx]
        if any(vals):
            rows.append(dict(zip(headers, vals)))
    return hr, headers, rows


def parse_money(value):
    """Parse currency-like text values into signed floats."""
    text = sval(value)
    if not text:
        return 0.0

    neg = False
    if text.startswith('(') and text.endswith(')'):
        neg = True
        text = text[1:-1]

    text = text.replace('$', '').replace(',', '').replace(' ', '')
    upper = text.upper()
    if upper.endswith('CR'):
        text = text[:-2]
    if upper.endswith('DR'):
        neg = True
        text = text[:-2]

    try:
        num = float(text)
    except ValueError:
        return 0.0

    if neg:
        return -abs(num)
    return num


def category_counts(rows):
    """Count rows by their detected exception or status category."""
    keys = ['Error Category', 'Category', 'Exception Category', 'Status']
    out = Counter()
    for row in rows:
        cat = ''
        for k in keys:
            if row.get(k):
                cat = row[k]
                break
        if cat:
            out[cat] += 1
    return out


def category_exposure(rows):
    """Aggregate exposure totals by detected exception category."""
    cat_keys = ['Error Category', 'Category', 'Exception Category', 'Status']
    amt_keys = ['Amount', 'Dollar Exposure', 'Exposure']
    out = Counter()
    for row in rows:
        cat = ''
        for k in cat_keys:
            if row.get(k):
                cat = row[k]
                break
        if not cat:
            continue

        amount = 0.0
        for k in amt_keys:
            if row.get(k):
                amount = parse_money(row[k])
                break

        out[cat] += amount
    return out


def print_summary_diff(sa, sb):
    """Print key-by-key differences between two summary sections."""
    keys = sorted(set(sa) | set(sb))
    diffs = []
    for k in keys:
        av = sa.get(k, '')
        bv = sb.get(k, '')
        if av != bv:
            diffs.append((k, av, bv))

    print(f'SUMMARY_DIFF_COUNT: {len(diffs)}')
    for k, av, bv in diffs:
        print(f'  - {k}: A={av!r} | B={bv!r}')


def main():
    """Compare workbook structure, summary values, and category deltas."""
    args = parse_args()
    a_path = Path(args.baseline).expanduser().resolve()
    b_path = Path(args.candidate).expanduser().resolve()

    if not a_path.exists() or not b_path.exists():
        print('ERROR: workbook path missing')
        print(f'  A exists: {a_path.exists()} => {a_path}')
        print(f'  B exists: {b_path.exists()} => {b_path}')
        raise SystemExit(1)

    wa = load_workbook(a_path, data_only=True)
    wb = load_workbook(b_path, data_only=True)

    print('WORKBOOK_A:', a_path)
    print('WORKBOOK_B:', b_path)
    print('')

    print('SHEETS_A:', wa.sheetnames)
    print('SHEETS_B:', wb.sheetnames)
    print('')

    if 'Summary' in wa.sheetnames and 'Summary' in wb.sheetnames:
        sa = read_summary(wa['Summary'])
        sb = read_summary(wb['Summary'])
        print_summary_diff(sa, sb)
    else:
        print('SUMMARY: missing in one of workbooks')
    print('')

    if (
        ERROR_REPORT_SHEET in wa.sheetnames
        and ERROR_REPORT_SHEET in wb.sheetnames
    ):
        hra, hda, ra = read_table(wa[ERROR_REPORT_SHEET])
        hrb, hdb, rb = read_table(wb[ERROR_REPORT_SHEET])

        print('ERROR_REPORT:')
        print(f'  A header_row={hra}, rows={len(ra)}, cols={len(hda)}')
        print(f'  B header_row={hrb}, rows={len(rb)}, cols={len(hdb)}')

        ca = category_counts(ra)
        cb = category_counts(rb)
        cats = sorted(set(ca) | set(cb))
        print('  CATEGORY_COUNT_DELTA:')
        for cat in cats:
            av = ca.get(cat, 0)
            bv = cb.get(cat, 0)
            if av != bv:
                print(f'    - {cat}: A={av}, B={bv}, delta={bv - av}')

        da = category_exposure(ra)
        db = category_exposure(rb)
        print('  CATEGORY_EXPOSURE_DELTA:')
        for cat in cats:
            av = round(da.get(cat, 0.0), 2)
            bv = round(db.get(cat, 0.0), 2)
            if av != bv:
                print(
                    f'    - {cat}: A={av}, B={bv}, '
                    f'delta={round(bv - av, 2)}'
                )
    else:
        print('ERROR_REPORT: missing in one of workbooks')

    print('')
    for sheet in REVIEW_SHEETS:
        if sheet in wa.sheetnames and sheet in wb.sheetnames:
            a = wa[sheet]
            b = wb[sheet]
            print(
                f'{sheet}: A rows={a.max_row}, B rows={b.max_row}, '
                f'A cols={a.max_column}, B cols={b.max_column}'
            )


if __name__ == '__main__':
    main()
