"""Shared workbook comparison helpers and CLI routines."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

ERROR_REPORT_SHEET = 'Error Report'
REVIEW_SHEETS = [
    'Mapping and Assumptions',
    'Probable Matches Review',
    'Reviewer Observations',
]
CATEGORY_KEYS = ['Error Category', 'Category', 'Exception Category', 'Status']
AMOUNT_KEYS = ['Amount', 'Dollar Exposure', 'Exposure']


def build_parser(description: str) -> argparse.ArgumentParser:
    """Build a baseline/candidate argument parser."""
    parser = argparse.ArgumentParser(description=description)
    parser.add_argument(
        '--baseline', required=True, help='Path to the baseline workbook.'
    )
    parser.add_argument(
        '--candidate', required=True, help='Path to the candidate workbook.'
    )
    return parser


def normalize_value(value):
    """Return a stable string form for workbook values."""
    if value is None:
        return ''
    if isinstance(value, float):
        return ('%.10f' % value).rstrip('0').rstrip('.')
    return str(value).strip()


def parse_money(value):
    """Parse workbook currency values into floats when possible."""
    text = normalize_value(value)
    if not text:
        return None

    negative = False
    if text.startswith('(') and text.endswith(')'):
        negative = True
        text = text[1:-1]

    text = text.replace('$', '').replace(',', '').replace(' ', '')
    upper = text.upper()
    if upper.endswith('CR'):
        text = text[:-2]
    if upper.endswith('DR'):
        negative = True
        text = text[:-2]

    try:
        amount = float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None

    return -abs(amount) if negative else amount


def resolve_workbook_paths(baseline: str, candidate: str):
    """Resolve workbook paths and fail fast if either is missing."""
    baseline_path = Path(baseline).expanduser().resolve()
    candidate_path = Path(candidate).expanduser().resolve()

    if not baseline_path.exists() or not candidate_path.exists():
        print('ERROR: workbook path missing')
        print(f'  A exists: {baseline_path.exists()} => {baseline_path}')
        print(f'  B exists: {candidate_path.exists()} => {candidate_path}')
        raise SystemExit(1)

    return baseline_path, candidate_path


def read_summary(worksheet, max_rows: int = 300):
    """Read the top summary key-value pairs from a worksheet."""
    summary = {}
    for row in range(1, min(worksheet.max_row, max_rows) + 1):
        key = normalize_value(worksheet.cell(row, 1).value)
        value = normalize_value(worksheet.cell(row, 2).value)
        if key:
            summary[key] = value
    return summary


def detect_header_row(worksheet, max_scan: int = 12, max_columns: int = 100):
    """Find the most likely header row by counting non-empty cells."""
    best_count = -1
    best_row = 1
    for row in range(1, min(worksheet.max_row, max_scan) + 1):
        non_empty = 0
        for column in range(1, min(worksheet.max_column, max_columns) + 1):
            if normalize_value(worksheet.cell(row, column).value):
                non_empty += 1
        if non_empty > best_count:
            best_count = non_empty
            best_row = row
    return best_row


def read_table(worksheet, max_columns: int = 100):
    """Read a worksheet into header metadata and row dictionaries."""
    header_row = detect_header_row(worksheet, max_columns=max_columns)
    columns = []
    headers = []
    for column in range(1, min(worksheet.max_column, max_columns) + 1):
        header = normalize_value(worksheet.cell(header_row, column).value)
        if header:
            columns.append(column)
            headers.append(header)

    rows = []
    for row in range(header_row + 1, worksheet.max_row + 1):
        values = [normalize_value(worksheet.cell(row, column).value) for column in columns]
        if any(values):
            rows.append(dict(zip(headers, values)))

    return header_row, headers, rows


def get_category(row):
    """Extract the category-like field from a workbook row."""
    for key in CATEGORY_KEYS:
        if row.get(key):
            return row[key]
    return ''


def get_amount(row):
    """Extract the monetary exposure field from a workbook row."""
    for key in AMOUNT_KEYS:
        if row.get(key):
            return parse_money(row[key]) or 0.0
    return 0.0


def read_error_report_rows(path: Path, mapping: dict[str, int]):
    """Read error report rows using a fixed column mapping."""
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook[ERROR_REPORT_SHEET]
    rows = []
    for row in range(2, worksheet.max_row + 1):
        source = normalize_value(worksheet.cell(row, mapping['source']).value)
        line = normalize_value(worksheet.cell(row, mapping['line']).value)
        category = normalize_value(worksheet.cell(row, mapping['category']).value)
        message = normalize_value(worksheet.cell(row, mapping['message']).value)
        amount = parse_money(worksheet.cell(row, mapping['amount']).value)
        date = None
        if 'date' in mapping:
            date = normalize_value(worksheet.cell(row, mapping['date']).value)

        if not source and not line and not category:
            continue

        rows.append(
            {
                'key': (source.lower(), line),
                'source': source,
                'line': line,
                'date': date,
                'amount': amount,
                'cat': category,
                'msg': message,
                'row': row,
            }
        )
    return rows


def run_parity(baseline: str, candidate: str):
    """Compare workbook structure, summary values, and category deltas."""
    baseline_path, candidate_path = resolve_workbook_paths(baseline, candidate)

    workbook_a = load_workbook(baseline_path, data_only=True)
    workbook_b = load_workbook(candidate_path, data_only=True)

    print('WORKBOOK_A:', baseline_path)
    print('WORKBOOK_B:', candidate_path)
    print('')
    print('SHEETS_A:', workbook_a.sheetnames)
    print('SHEETS_B:', workbook_b.sheetnames)
    print('')

    if 'Summary' in workbook_a.sheetnames and 'Summary' in workbook_b.sheetnames:
        summary_a = read_summary(workbook_a['Summary'], max_rows=250)
        summary_b = read_summary(workbook_b['Summary'], max_rows=250)
        keys = sorted(set(summary_a) | set(summary_b))
        diffs = []
        for key in keys:
            value_a = summary_a.get(key, '')
            value_b = summary_b.get(key, '')
            if value_a != value_b:
                diffs.append((key, value_a, value_b))

        print(f'SUMMARY_DIFF_COUNT: {len(diffs)}')
        for key, value_a, value_b in diffs:
            print(f'  - {key}: A={value_a!r} | B={value_b!r}')
    else:
        print('SUMMARY: missing in one of workbooks')
    print('')

    if ERROR_REPORT_SHEET in workbook_a.sheetnames and ERROR_REPORT_SHEET in workbook_b.sheetnames:
        header_a, headers_a, rows_a = read_table(workbook_a[ERROR_REPORT_SHEET], max_columns=80)
        header_b, headers_b, rows_b = read_table(workbook_b[ERROR_REPORT_SHEET], max_columns=80)
        print('ERROR_REPORT:')
        print(f'  A header_row={header_a}, rows={len(rows_a)}, cols={len(headers_a)}')
        print(f'  B header_row={header_b}, rows={len(rows_b)}, cols={len(headers_b)}')

        counts_a = Counter(get_category(row) for row in rows_a if get_category(row))
        counts_b = Counter(get_category(row) for row in rows_b if get_category(row))
        categories = sorted(set(counts_a) | set(counts_b))
        print('  CATEGORY_COUNT_DELTA:')
        for category in categories:
            count_a = counts_a.get(category, 0)
            count_b = counts_b.get(category, 0)
            if count_a != count_b:
                print(f'    - {category}: A={count_a}, B={count_b}, delta={count_b - count_a}')

        exposure_a = Counter()
        exposure_b = Counter()
        for row in rows_a:
            category = get_category(row)
            if category:
                exposure_a[category] += get_amount(row)
        for row in rows_b:
            category = get_category(row)
            if category:
                exposure_b[category] += get_amount(row)

        print('  CATEGORY_EXPOSURE_DELTA:')
        for category in categories:
            value_a = round(exposure_a.get(category, 0.0), 2)
            value_b = round(exposure_b.get(category, 0.0), 2)
            if value_a != value_b:
                print(
                    f'    - {category}: A={value_a}, B={value_b}, '
                    f'delta={round(value_b - value_a, 2)}'
                )
    else:
        print('ERROR_REPORT: missing in one of workbooks')

    print('')
    for sheet in REVIEW_SHEETS:
        if sheet in workbook_a.sheetnames and sheet in workbook_b.sheetnames:
            review_a = workbook_a[sheet]
            review_b = workbook_b[sheet]
            print(
                f'{sheet}: A rows={review_a.max_row}, B rows={review_b.max_row}, '
                f'A cols={review_a.max_column}, B cols={review_b.max_column}'
            )


def run_row_compare(baseline: str, candidate: str):
    """Compare workbook summaries, error categories, and row-level deltas."""
    baseline_path, candidate_path = resolve_workbook_paths(baseline, candidate)

    workbook_a = load_workbook(baseline_path, data_only=True)
    workbook_b = load_workbook(candidate_path, data_only=True)

    print('A', baseline_path)
    print('B', candidate_path)
    print('SHEETS_A', workbook_a.sheetnames)
    print('SHEETS_B', workbook_b.sheetnames)

    if 'Summary' in workbook_a.sheetnames and 'Summary' in workbook_b.sheetnames:
        summary_a = read_summary(workbook_a['Summary'])
        summary_b = read_summary(workbook_b['Summary'])
        keys = sorted(set(summary_a) | set(summary_b))
        diffs = [
            (key, summary_a.get(key, ''), summary_b.get(key, ''))
            for key in keys
            if summary_a.get(key, '') != summary_b.get(key, '')
        ]
        print('SUMMARY_DIFF_COUNT', len(diffs))
        for key, value_a, value_b in diffs[:60]:
            print('SUMMARY_DIFF', key, '|A=', value_a, '|B=', value_b)

    if ERROR_REPORT_SHEET in workbook_a.sheetnames and ERROR_REPORT_SHEET in workbook_b.sheetnames:
        header_a, headers_a, rows_a = read_table(workbook_a[ERROR_REPORT_SHEET])
        header_b, headers_b, rows_b = read_table(workbook_b[ERROR_REPORT_SHEET])
        print('ERROR_REPORT A_header', header_a, 'A_rows', len(rows_a), 'A_cols', len(headers_a))
        print('ERROR_REPORT B_header', header_b, 'B_rows', len(rows_b), 'B_cols', len(headers_b))

        counts_a = Counter(get_category(row) for row in rows_a if get_category(row))
        counts_b = Counter(get_category(row) for row in rows_b if get_category(row))
        categories = sorted(set(counts_a) | set(counts_b))
        print('CATEGORY_COUNT_DELTAS')
        for category in categories:
            if counts_a.get(category, 0) != counts_b.get(category, 0):
                print(
                    'CAT',
                    category,
                    'A',
                    counts_a.get(category, 0),
                    'B',
                    counts_b.get(category, 0),
                    'DELTA',
                    counts_b.get(category, 0) - counts_a.get(category, 0),
                )

        exposure_a = Counter()
        exposure_b = Counter()
        for row in rows_a:
            category = get_category(row)
            if category:
                exposure_a[category] += get_amount(row)
        for row in rows_b:
            category = get_category(row)
            if category:
                exposure_b[category] += get_amount(row)
        print('CATEGORY_EXPOSURE_DELTAS')
        for category in categories:
            value_a = round(exposure_a.get(category, 0.0), 2)
            value_b = round(exposure_b.get(category, 0.0), 2)
            if value_a != value_b:
                print('EXP', category, 'A', value_a, 'B', value_b, 'DELTA', round(value_b - value_a, 2))

    rows_a = read_error_report_rows(
        baseline_path,
        {'source': 1, 'line': 3, 'amount': 9, 'category': 10, 'message': 11},
    )
    rows_b = read_error_report_rows(
        candidate_path,
        {'source': 3, 'line': 5, 'amount': 7, 'category': 2, 'message': 8},
    )

    index_a = defaultdict(list)
    index_b = defaultdict(list)
    for row in rows_a:
        index_a[row['key']].append(row)
    for row in rows_b:
        index_b[row['key']].append(row)

    common = set(index_a) & set(index_b)
    only_a = set(index_a) - set(index_b)
    only_b = set(index_b) - set(index_a)

    print('KEY_COVERAGE A', len(index_a), 'B', len(index_b), 'COMMON', len(common), 'ONLY_A', len(only_a), 'ONLY_B', len(only_b))

    mismatch = Counter()
    for key in common:
        left = index_a[key][0]['cat']
        right = index_b[key][0]['cat']
        if left != right:
            mismatch[(left, right)] += 1

    print('TOP_MISMATCH_PAIRS')
    for (left, right), count in mismatch.most_common(20):
        print('PAIR', left, '->', right, 'COUNT', count)

    print('SAMPLE_ONLY_B')
    for key in sorted(only_b)[:30]:
        row = index_b[key][0]
        print(
            'ONLY_B',
            row['source'],
            'line',
            row['line'],
            'cat',
            row['cat'],
            'amt',
            row['amount'],
            'msg',
            row['msg'][:110],
        )


def run_second_pass(baseline: str, candidate: str):
    """Print category, coverage, and exposure diagnostics."""
    baseline_path, candidate_path = resolve_workbook_paths(baseline, candidate)
    rows_a = read_error_report_rows(
        baseline_path,
        {'source': 1, 'line': 3, 'date': 4, 'amount': 9, 'category': 10, 'message': 11},
    )
    rows_b = read_error_report_rows(
        candidate_path,
        {'source': 3, 'line': 5, 'date': 6, 'amount': 7, 'category': 2, 'message': 8},
    )

    index_a = defaultdict(list)
    index_b = defaultdict(list)
    for row in rows_a:
        index_a[row['key']].append(row)
    for row in rows_b:
        index_b[row['key']].append(row)

    common = set(index_a).intersection(index_b)
    only_a = set(index_a) - set(index_b)
    only_b = set(index_b) - set(index_a)

    print('KEY COVERAGE')
    print(
        f'A keys: {len(index_a)} B keys: {len(index_b)} Common: {len(common)} '
        f'OnlyA: {len(only_a)} OnlyB: {len(only_b)}'
    )

    mismatch = Counter()
    matches = Counter()
    examples = []
    for key in sorted(common):
        row_a = index_a[key][0]
        row_b = index_b[key][0]
        category_a = row_a['cat']
        category_b = row_b['cat']
        if category_a == category_b:
            matches[category_a] += 1
        else:
            mismatch[(category_a, category_b)] += 1
            delta = None
            if row_a['amount'] is not None and row_b['amount'] is not None:
                delta = round(row_b['amount'] - row_a['amount'], 2)
            examples.append((abs(row_a['amount'] or 0), row_a, row_b, delta))

    print('\nTOP CATEGORY MISMATCH PAIRS (A -> B)')
    for (category_a, category_b), count in mismatch.most_common(15):
        print(f'{count:4d} | {category_a} -> {category_b}')

    print('\nTOP CATEGORY MATCHES')
    for category, count in matches.most_common(10):
        print(f'{count:4d} | {category}')

    print(f'\nONLY IN A (baseline-only keys): {len(only_a)}')
    print(f'ONLY IN B (app-only keys): {len(only_b)}')

    print('\nSAMPLE ONLY-IN-B KEYS (first 30)')
    for key in sorted(only_b)[:30]:
        row = index_b[key][0]
        print(
            f"{row['source']} line {row['line']} | cat={row['cat']} | "
            f"amount={row['amount']} | "
            f"msg={row['msg'][:90]}"
        )

    print('\nTOP 50 HIGH-IMPACT MISMATCH EXAMPLES (by abs baseline amount)')
    examples.sort(reverse=True, key=lambda item: item[0])
    for _, row_a, row_b, delta in examples[:50]:
        print(
            f"{row_a['source']} line {row_a['line']} | A:{row_a['cat']} "
            f"amt={row_a['amount']} -> "
            f"B:{row_b['cat']} amt={row_b['amount']} | dAmt={delta} | "
            f"Amsg={row_a['msg'][:70]} | Bmsg={row_b['msg'][:70]}"
        )

    sum_a = Counter()
    sum_b = Counter()
    for row in rows_a:
        if row['amount'] is not None:
            sum_a[row['cat']] += row['amount']
    for row in rows_b:
        if row['amount'] is not None:
            sum_b[row['cat']] += row['amount']

    print('\nCATEGORY EXPOSURE DELTA (B - A)')
    for category in sorted(set(sum_a) | set(sum_b)):
        delta = round(sum_b[category] - sum_a[category], 2)
        if abs(delta) > 0.01:
            print(f"{category}: {delta:+,.2f} (A={sum_a[category]:,.2f} B={sum_b[category]:,.2f})")
